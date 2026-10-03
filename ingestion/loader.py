"""Read user-provided original documents without downloading a dataset."""
from dataclasses import dataclass
import hashlib
import os
from pathlib import Path

from core.config_loader import ROOT

SUPPORTED = {'.md', '.txt', '.pdf', '.docx', '.html'}


@dataclass
class Document:
    doc_id: str
    title: str
    text: str
    source: str
    content_hash: str


def source_directory(source_id):
    if source_id == 'project':
        directory = Path(os.getenv('RAG_DOCUMENTS_DIR', './data/documents'))
        return directory if directory.is_absolute() else ROOT / directory
    if not source_id.startswith('upload_'):
        raise ValueError('Unknown document source')
    return ROOT / 'data' / 'uploads' / source_id


def document_paths(source_id):
    directory = source_directory(source_id)
    if not directory.exists():
        return []
    return sorted(path for path in directory.rglob('*') if path.is_file() and not path.is_symlink() and path.suffix.lower() in SUPPORTED)


def source_revision(paths, directory):
    hasher = hashlib.sha256()
    for path in paths:
        hasher.update(str(path.relative_to(directory)).encode())
        with path.open('rb') as file:
            for block in iter(lambda: file.read(1024 * 1024), b''):
                hasher.update(block)
    return hasher.hexdigest()


def load_document(path, directory):
    raw = path.read_bytes()
    extension = path.suffix.lower()
    if extension == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError(f'{path.name}: encrypted PDFs are not supported')
        text = '\n\n'.join(page.extract_text() or '' for page in reader.pages)
    elif extension == '.docx':
        from docx import Document as WordDocument
        document = WordDocument(path)
        text = '\n\n'.join(paragraph.text for paragraph in document.paragraphs)
        for table in document.tables:
            text += '\n\n' + '\n'.join(' | '.join(cell.text for cell in row.cells) for row in table.rows)
    else:
        text = raw.decode('utf-8-sig')
    if not text.strip():
        raise ValueError(f'{path.name}: no extractable text; scanned PDFs require OCR before ingestion')
    relative = str(path.relative_to(directory))
    content_hash = hashlib.sha256(raw).hexdigest()
    doc_id = hashlib.sha256((relative + ':' + content_hash).encode()).hexdigest()[:24]
    title = next((line.lstrip('# ').strip() for line in text.splitlines() if line.startswith('# ')), path.stem)
    return Document(doc_id, title, text, relative, content_hash)
